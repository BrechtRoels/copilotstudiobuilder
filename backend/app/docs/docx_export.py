"""Doc -> Word .docx. Diagrams come in as PNGs rendered by the browser (Mermaid)."""
from __future__ import annotations

import base64
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from .builder import Doc

ORANGE = RGBColor(0xD0, 0x4A, 0x02)
GREY = RGBColor(0x5C, 0x5C, 0x5C)


def _heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = ORANGE if level <= 1 else RGBColor(0x2D, 0x2D, 0x2D)
        r.font.name = "Arial"


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 2"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v or "—"
    doc.add_paragraph()


def _kv(doc: Document, pairs: list[tuple[str, str]]) -> None:
    for k, v in pairs:
        p = doc.add_paragraph()
        r = p.add_run(f"{k}: ")
        r.bold = True
        p.add_run(v or "—")
        p.paragraph_format.space_after = Pt(2)


def _image(doc: Document, images: dict[str, str], key: str) -> None:
    data = images.get(key)
    if not data:
        return
    raw = base64.b64decode(data.split(",", 1)[-1])
    doc.add_picture(io.BytesIO(raw), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _pre(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)


def render_docx(d: Doc, images: dict[str, str] | None = None) -> bytes:
    images = images or {}
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    title = doc.add_heading(d.title, level=0)
    for r in title.runs:
        r.font.color.rgb = ORANGE
    meta = doc.add_paragraph(f"Agent documentation · generated {d.generatedAt} · {d.orchestration} orchestration"
                             + (f" · model {d.model}" if d.model else ""))
    meta.runs[0].font.color.rgb = GREY
    _table(doc, ["Topics", "Child agents", "Tools", "Agent flows", "Knowledge", "Global variables"],
           [[str(d.counts[k]) for k in ("topics", "childAgents", "tools", "flows", "knowledge", "variables")]])

    _heading(doc, "1. Overview", 1)
    if d.overview:
        doc.add_paragraph(d.overview)
    if d.instructions:
        _heading(doc, "Agent instructions", 2)
        _pre(doc, d.instructions)
    if d.starters:
        _kv(doc, [("Conversation starters", " · ".join(d.starters))])

    _heading(doc, "2. Architecture", 1)
    if "agent-map" in images:
        _image(doc, images, "agent-map")
    else:
        doc.add_paragraph("See the HTML export for the interactive architecture diagram.")

    n = 3
    if d.childAgents:
        _heading(doc, f"{n}. Child agents", 1)
        n += 1
        for a in d.childAgents:
            _heading(doc, a["name"], 2)
            if a["summary"]:
                doc.add_paragraph(a["summary"])
            _kv(doc, [("Routing description", a["description"]), ("Inputs", ", ".join(a["inputs"])),
                      ("Outputs", ", ".join(a["outputs"]))])
            if a["instructions"]:
                _pre(doc, a["instructions"])

    _heading(doc, f"{n}. Topics", 1)
    n += 1
    _table(doc, ["Topic", "Trigger", "Routing", "Nodes"],
           [[t.name, t.trigger, t.description or " · ".join(t.phrases[:3]), str(t.nodeCount)] for t in d.topics])
    for t in d.topics:
        _heading(doc, t.name, 2)
        if t.summary:
            doc.add_paragraph(t.summary)
        _kv(doc, [("Trigger", t.trigger), ("Model description", t.description),
                  ("Trigger phrases", " · ".join(t.phrases)), ("Variables", ", ".join(t.variables)),
                  ("Redirects to", ", ".join(t.redirects)), ("Flows called", ", ".join(t.flows))])
        _image(doc, images, t.diagram)

    if d.tools:
        _heading(doc, f"{n}. Tools", 1)
        n += 1
        _table(doc, ["Tool", "Type", "Target", "Inputs", "Description"],
               [[t["name"], t["type"], t["target"], ", ".join(t["inputs"]), t["description"]] for t in d.tools])

    if d.flows:
        _heading(doc, f"{n}. Agent flows", 1)
        n += 1
        for f in d.flows:
            _heading(doc, f.name, 2)
            if f.summary:
                doc.add_paragraph(f.summary)
            _kv(doc, [("Trigger", f.trigger), ("Inputs", ", ".join(f.inputs)), ("Outputs", ", ".join(f.outputs)),
                      ("Connectors", ", ".join(f.connectors)),
                      ("Error handling", "Yes" if f.errorHandling else "No failure path")])
            _image(doc, images, f.diagram)
            if f.steps:
                for s in f.steps:
                    doc.add_paragraph(s, style="List Bullet")

    if d.knowledge or d.variables or d.connections:
        _heading(doc, f"{n}. Knowledge, variables & connections", 1)
        n += 1
        if d.knowledge:
            _table(doc, ["Knowledge source", "Type", "Used by", "Location"],
                   [[k["name"], k["type"], k["owner"], k["detail"]] for k in d.knowledge])
        if d.variables:
            _table(doc, ["Global variable", "AI visibility", "Description"],
                   [[v["name"], v["visibility"], v["description"]] for v in d.variables])
        if d.connections:
            _table(doc, ["Connection reference"], [[c] for c in d.connections])

    _heading(doc, f"{n}. Quality findings", 1)
    if d.findings:
        order = {"error": 0, "warning": 1, "info": 2}
        _table(doc, ["Severity", "Component", "Finding"],
               [[f.severity, f.component, f.message] for f in sorted(d.findings, key=lambda f: order[f.severity])])
    else:
        doc.add_paragraph("No findings.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
